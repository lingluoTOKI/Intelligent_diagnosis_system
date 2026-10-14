"""Generate four formal course-design reports from verified project evidence."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "course_design"
FIG = ROOT / "docs" / "evidence" / "report_figures"
if not FIG.is_dir():
    FIG = ROOT / "runs" / "report_figures"
TITLE = "基于YOLO与大语言模型的眼科图像分类辅助诊断系统"
MEMBERS = "欧阳宸（230722105）  汪晓蕾（231613212）"
DATE = "2026年10月"


def shade(cell, fill="1F4E78"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), color)
        borders.append(tag)
    tc_pr.append(borders)


def set_run(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text="", bold_lead=None, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run(r)
    else:
        r = p.add_run(text)
        set_run(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r, 16 if level == 1 else 14, True)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_marker = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header_marker)
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = h
        shade(c)
        border(c)
        for r in c.paragraphs[0].runs:
            set_run(r, 10.5, True)
            r.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            border(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.line_spacing = 1.25
                for r in p.runs:
                    set_run(r, 10.5)
    if widths:
        available = 21 - 2 * 2.54
        widths = [width * available / sum(widths) for width in widths]
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def figure(doc, file, caption):
    path = FIG / file
    if path.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(14.8))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        set_run(r, 10.5)


def setup(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.54)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header.add_run("智能软件系统课程设计")
    set_run(r, 9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("眼科图像分类辅助诊断系统")
    set_run(r, 9)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(5)
    r = p.add_run("智能软件系统课程设计")
    set_run(r, 22, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(2)
    r = p.add_run(title); set_run(r, 20, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(4)
    r = p.add_run(TITLE); set_run(r, 15)
    for label, value in (("组员", MEMBERS), ("提交日期", DATE)):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Cm(0.8)
        r = p.add_run(f"{label}：{value}"); set_run(r, 13)
    doc.add_page_break()
    return doc


def references(doc):
    p = add_heading(doc, "参考文献", 1)
    p.paragraph_format.page_break_before = True
    items = [
        "[1] ODIR-2019. Dataset ODIR-5K [EB/OL]. https://odir2019.grand-challenge.org/dataset/, accessed 2026-10-14. 数据集说明页面，不计作核心研究论文。",
        "[2] D. S. W. Ting et al. Artificial intelligence and deep learning in ophthalmology. British Journal of Ophthalmology, 2019, 103(2): 167-175. doi:10.1136/bjophthalmol-2018-313173.",
        "[3] V. Gulshan et al. Development and validation of a deep learning algorithm for detection of diabetic retinopathy. JAMA, 2016.",
        "[4] D. S. W. Ting et al. Development and validation of a deep learning system for diabetic retinopathy and related eye diseases. JAMA, 2017.",
        "[5] G. Litjens et al. A survey on deep learning in medical image analysis. Medical Image Analysis, 2017.",
        "[6] X. Zhang et al. AKConv: Convolutional Kernel with Arbitrary Sampled Shapes and Arbitrary Number of Parameters. arXiv:2311.11587v2, 2023. https://arxiv.org/abs/2311.11587v2.",
        "[7] Ultralytics. YOLO Classification Documentation and Source Code. https://docs.ultralytics.com, accessed 2026-10-13.",
        "[8] X. Liu et al. A comparison of deep learning performance against health-care professionals in detecting diseases from medical imaging. The Lancet Digital Health, 2019.",
        "[9] DeepSeek-AI. DeepSeek V3 Technical Report. arXiv:2412.19437, 2024.",
        "[10] Y. Zhou et al. A foundation model for generalizable disease detection from retinal images. Nature, 2023, 622: 156-163. doi:10.1038/s41586-023-06555-x.",
    ]
    for item in items:
        add_para(doc, item, indent=False)


def literature():
    doc = setup("文献分析报告")
    add_heading(doc, "1 研究背景与意义")
    add_para(doc, "眼底图像能够反映视网膜、视神经和黄斑区域的病变线索。临床筛查通常需要专业设备和人员判读，而基层或远程场景容易受资源限制。本项目以八类眼科图像分类为目标，提供图像导入、分类概率、历史记录和辅助说明，定位为课程设计中的辅助信息系统，不输出临床诊断结论。")
    add_para(doc, "项目的价值在于把分类模型、桌面交互、网页服务与语言模型说明组织在同一可运行流程中。模型只给出训练类别范围内的统计预测，建议模块保留来源标识和风险提示，避免把生成文本误当作诊疗意见。")
    add_heading(doc, "2 国内外研究现状")
    add_para(doc, "公开眼底数据集推动了自动筛查研究。ODIR 5K 提供双眼彩色眼底照片及多个疾病标签，为多类别研究提供了基础。深度学习方法已经在糖尿病视网膜病变等任务中显示出较好的筛查能力，但跨设备、类别不均衡、样本独立性和临床外部验证仍是决定结果可信度的关键因素。")
    add_para(doc, "轻量级卷积网络适合有限显存和边缘部署。YOLO 分类模型可复用成熟训练、增强和导出工具链。AKConv 通过可学习采样位置扩展卷积的几何适应性，但其有效性必须由相同数据和独立验证下的消融实验支持，不能仅依据代码中存在模块或目录名称作结论。")
    add_para(doc, "大语言模型可将类别名称、概率和限制条件组织为可读说明，适合辅助交互；其输出会受提示、上下文和服务状态影响，因此应显示来源、限制和失败信息。")
    add_para(doc, "Zhou等人的RETFound研究以未标注视网膜图像的自监督预训练为基础，再适配具体疾病检测任务，说明领域预训练是值得比较的技术路线[10]。本项目尚未训练或评估RETFound，不能将其公开结果与本项目受污染划分上的指标直接排序；它仅作为后续独立实验的候选基准。")
    add_heading(doc, "3 方法对比与评述")
    add_table(doc, ["方向", "代表工作", "优势", "主要局限", "与本项目关系"], [
        ["眼底数据集", "ODIR 5K", "类别覆盖广，含双眼图像", "多标签与患者级信息需谨慎处理", "作为历史来源线索"],
        ["深度学习筛查", "Gulshan 等", "端到端特征学习", "外部泛化依赖独立数据", "采用图像分类范式"],
        ["轻量网络", "YOLO 分类", "训练和导出工具完整", "类别不均衡会影响结果", "桌面与服务端推理基础"],
        ["动态卷积", "AKConv", "可适应不规则局部结构", "增加训练成本，需消融验证", "用于候选结构对比"],
        ["语言模型", "DeepSeek", "可生成结构化说明", "可能产生不可靠内容", "仅作辅助说明"],
    ], [2.3, 2.5, 3.4, 4.2, 3.4])
    add_heading(doc, "4 项目切入点与边界")
    add_para(doc, "本项目的切入点不是宣称替代医生，而是构建可追溯的图像分类辅助工作流：分类失败不写入成功历史；建议区说明结果来源；网页端和桌面端采用相同八类映射；服务器使用 ONNX Runtime 进行轻量推理。")
    add_para(doc, "数据审计发现训练集与验证集存在 140 个跨集合重复内容哈希。因此，已有验证指标只可作为历史训练记录，不可用于证明独立泛化性能。后续应基于原始元数据进行去重和患者级划分，再开展正式比较。")
    references(doc)
    return doc


def design():
    doc = setup("技术方案设计报告")
    add_heading(doc, "1 需求分析")
    add_para(doc, "用户需要能够导入单张或批量眼底图像，获得八类分类概率并保存可查询记录；需要在图像无效、模型未加载、接口不可用时看到明确原因。教学演示还需要摄像头采集、语音输入、朗读、辅助说明和网页服务入口。")
    add_table(doc, ["类别", "需求", "验收方式"], [
        ["功能", "单图、批量、历史查询、导出、辅助说明和多轮问答", "测试用例与界面操作记录"],
        ["性能", "本机 GPU 推理与服务器 CPU 推理可用", "服务健康接口和实际单图计时"],
        ["可靠性", "异常图像和模型失败不产生伪成功历史", "自动化测试与异常样例"],
        ["安全边界", "密钥不写入仓库，医疗结果显示辅助提示", "忽略规则和页面提示"],
        ["硬件", "本机摄像头、麦克风、扬声器；开发板可选", "开发板未实机验收时明确标注"],
    ], [2.2, 8.5, 5.1])
    add_heading(doc, "2 总体架构")
    add_para(doc, "桌面端以 PyQt5 工作台为入口，负责图像导入、摄像头、展示和本机偏好。分类模块加载 YOLO 分类权重，历史模块使用 SQLite 保存分类记录。DeepSeek 接口在配置可用时返回在线说明，未配置时返回本地参考内容。网页端以 FastAPI 提供上传、批量、历史、问答和语音接口，ONNX Runtime 在服务器 CPU 上推理。")
    add_table(doc, ["层次", "组成", "责任"], [
        ["交互层", "PyQt5 桌面端、响应式网页", "图像输入、结果展示、操作反馈"],
        ["业务层", "分类、问答、语音、批量任务", "流程编排、异常返回、缓存"],
        ["数据层", "SQLite、配置、运行日志", "历史、批次、偏好与可追溯记录"],
        ["推理层", "YOLO 分类权重、ONNX Runtime", "八类概率输出"],
        ["部署层", "启动器、systemd、部署包", "桌面启动与服务常驻"],
    ], [2.2, 5.0, 8.6])
    add_heading(doc, "3 数据与控制流程")
    add_para(doc, "单图流程为：选择或采集图像，检查格式和尺寸，执行预处理与模型推理，返回类别和概率，成功时保存历史记录，失败时返回原因且不保存成功记录。批量流程为逐项分类，登记批次号，成功项和失败项分别记录，历史页面可按日期、类别和批次筛选。")
    add_para(doc, "辅助建议流程为：分类结果或用户问题进入提示构造，在线服务可用时请求语言模型；不可用时返回本地参考说明。前端将来源和“辅助参考”边界一并展示。朗读优先通过服务端生成自然语音 MP3，并按文本缓存；失败时降级为浏览器本机朗读。")
    add_heading(doc, "4 技术选型与接口")
    add_table(doc, ["技术", "选型依据", "实际用途"], [
        ["Python 与 PyQt5", "桌面端开发和现有工程基础", "桌面工作台"],
        ["Ultralytics YOLO 分类", "训练、验证和导出流程完整", "本机模型推理"],
        ["ONNX Runtime", "低内存、适合 2GB 服务器", "服务端 CPU 推理"],
        ["FastAPI", "接口清晰、便于静态网页协作", "网页服务"],
        ["SQLite", "单机演示部署简单", "历史与批次记录"],
        ["Edge TTS", "中文自然音色和 MP3 缓存", "网页端朗读"],
    ], [3.0, 5.8, 7.0])
    add_para(doc, "核心接口包括 GET /api/health、POST /api/predict、POST /api/predict/batch、GET /api/history、POST /api/chat 和 POST /api/tts。服务端对上传文件类型和大小进行限制，模型不可用时返回 503 状态，而非伪造分类结果。")
    add_heading(doc, "5 部署与边界")
    add_para(doc, "服务器部署使用 ONNX 模型、requirements-server.txt 和 systemd 服务。服务器端已完成部署与运行；公网域名链路仍应以部署后健康接口和真实页面访问结果为准。开发板功能保留接口说明，但本项目验收时没有开发板，报告中标记为未实机验收。")
    references(doc)
    return doc


def model_report():
    doc = setup("模型分析报告")
    add_heading(doc, "1 基础模型与类别定义")
    add_para(doc, "部署任务为八类眼底图像分类：AMD、白内障、糖尿病视网膜病变、青光眼、高血压性视网膜病变、近视性黄斑病变、正常和其他。当前服务端权重为 ONNX 格式 best512.onnx，模型声明输入为 512×512；服务端以权重声明尺寸进行短边缩放和中心裁剪。")
    add_table(doc, ["项目", "已核实信息"], [
        ["部署骨干", "YOLO11n-cls 分类权重，原生部署权重不含 AKConv"],
        ["部署参数量", "1,448,536；与历史对比实验基线的 1,541,352 不同"],
        ["输入尺寸", "512×512（权重声明，服务端健康接口可查询）"],
        ["类别数", "8"],
        ["服务端运行时", "ONNX Runtime CPUExecutionProvider"],
        ["模型输出", "八类概率；用于辅助参考，不代表临床诊断准确率"],
    ], [4.0, 11.8])
    add_heading(doc, "2 训练与推理策略")
    add_para(doc, "历史训练记录使用 FP32、batch=32、workers=8、224 尺寸和 50 轮设置，采用分类交叉熵训练过程。该设置用于比较训练日志，不等同于当前部署权重的 512 推理预处理。部署侧必须以模型自身输入形状为准，避免将 224 导出模型与桌面推理混用。")
    add_para(doc, "推理模块首先校验 NCHW 三通道输入和八类输出，再完成归一化、模型执行与 softmax 判断。分类失败时向上层返回明确错误；业务层不保存成功历史，也不生成当前图像的伪报告。")
    add_heading(doc, "3 AKConv 适配与能力边界")
    add_para(doc, "工程中保留 AKConv 模块与候选训练记录。历史对比日志显示，AKConv 候选模型包含 6 个 AKConv 层，参数量为 1,487,772；基线模型参数量为 1,541,352。当前实际部署权重经结构核对 AKConv 层数为 0，因而系统不能宣称在线部署使用了 AKConv。")
    add_para(doc, "两套框架下的记录使用相同数据指纹，这是控制数据变量的正常条件，本身不能证明结果不独立。2026年10月14日进一步读取四个候选最佳检查点，以同一项目框架比较逐层模块与参数形状以及权重张量，发现两套框架中的基线模型结构和张量哈希分别相同，AKConv候选也分别相同。因此四条记录只能描述不同框架环境下的训练过程，不能作为四种不同网络架构的性能比较；哈希相同也不能单独用于断定训练日志被复制。")
    add_heading(doc, "4 资源与风险分析")
    add_table(doc, ["风险", "表现", "控制措施"], [
        ["数据泄漏", "训练与验证跨集合重复内容哈希 140 个", "去重、患者级划分、重训重测"],
        ["类别不均衡", "正常和糖尿病类样本占比较大", "报告每类指标，采用宏平均指标"],
        ["概率校准", "高概率不等于诊断正确", "界面显示辅助提示和限制"],
        ["语言模型幻觉", "建议文本可能不可靠", "显示来源，保留本地参考降级"],
        ["硬件差异", "服务器 CPU 延迟高于本机 GPU", "分别记录环境和测量值"],
    ], [3.1, 5.6, 7.1])
    add_heading(doc, "5 结论")
    add_para(doc, "当前模型可支撑课程演示中的八类图像辅助分类与服务器部署。严格有效性评价的前提仍是修复数据划分问题，并在独立样本上重新报告准确率、宏平均 F1、混淆矩阵和不确定性。")
    references(doc)
    return doc


def experiment():
    doc = setup("实验分析报告")
    add_heading(doc, "1 实验环境与数据构成")
    add_para(doc, "桌面端测试环境为 Windows PC，显卡为 NVIDIA GeForce RTX 4060 Laptop GPU；服务端为 Ubuntu 22.04、2 核 CPU、2GB 内存，使用 ONNX Runtime。历史数据目录统计显示训练集 5110 张，验证集 1163 张，八类分布不均衡。")
    add_table(doc, ["类别", "训练集", "验证集"], [
        ["AMD", 212, 35], ["白内障", 234, 55], ["糖尿病视网膜病变", 1286, 236],
        ["青光眼", 227, 61], ["高血压性视网膜病变", 102, 79], ["近视性黄斑病变", 185, 47],
        ["正常", 2298, 575], ["其他", 566, 75], ["合计", 5110, 1163],
    ], [7.0, 4.5, 4.5])
    figure(doc, "dataset_distribution.png", "图1 数据集类别分布")
    add_heading(doc, "2 数据审计与指标解释")
    add_para(doc, "审计发现训练集和验证集之间有 140 个重复内容的 SHA-256 哈希。该问题会使验证集结果受污染，因此本报告将已有训练曲线和 Top-1 数值作为历史实验记录，而不将其作为独立泛化性能结论。文件名前缀还存在 889 个潜在重叠线索，但因缺少原始患者元数据，不能直接把它解释为患者数量。")
    add_para(doc, "正式后续实验应使用去重且按患者划分的训练、验证和测试清单，报告准确率、宏平均精确率、宏平均召回率、宏平均 F1、每类混淆矩阵和置信区间。")
    add_heading(doc, "3 历史训练记录")
    add_table(doc, ["候选模型", "轮数", "历史最佳 Top-1", "参数量", "AKConv 层数", "说明"], [
        ["YOLO11n-cls", 50, "63.285%", "1,541,352", 0, "历史记录，受数据污染限制"],
        ["YOLO11n-cls + AKConv", 50, "63.199%", "1,487,772", 6, "历史记录，受数据污染限制"],
        ["YOLO26n-cls", 50, "63.285%", "1,541,352", 0, "与基线记录高度一致，非独立结论"],
        ["YOLO26n-cls + AKConv", 50, "63.199%", "1,487,772", 6, "与候选记录高度一致，非独立结论"],
    ], [3.8, 1.6, 3.0, 2.7, 2.2, 3.7])
    figure(doc, "curves_group_1.png", "图2 历史训练曲线（仅供训练过程复现）")
    figure(doc, "confusion_group_1.png", "图3 历史混淆矩阵（不作为独立验证结论）")
    add_heading(doc, "4 服务端与系统验证")
    add_para(doc, "服务端以 512 输入 ONNX 模型完成单图推理。已记录的 20 张同源图像对比中，桌面端 PyTorch 与 ONNX Runtime 的 Top-1 一致为 20/20，单类概率最大偏差约 0.000003；此验证说明导出推理口径一致，不代表临床或独立数据集准确率。服务器 CPU 单图推理记录为 192至204毫秒。")
    add_para(doc, "自动化测试覆盖分类失败不写成功历史、异常图像、模型加载失败、API 异常模拟、批量保存失败、历史过滤、导出、阅读字号和偏好恢复。摄像头、麦克风和扬声器需按人工测试记录确认；开发板因设备不在场，状态为未实机验收。")
    add_heading(doc, "5 结论与改进")
    add_para(doc, "系统已实现桌面和网页两种交互形式，并完成 ONNX 服务端部署。后续最优先工作是依据原始 ODIR 标签表完成数据来源复核、去重和患者级划分，重新训练并开展独立测试。只有完成该步骤后，模型比较与临床相关讨论才具有更可靠的基础。")
    references(doc)
    return doc


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    reports = {
        "文献分析报告.docx": literature(),
        "技术方案设计报告.docx": design(),
        "模型分析报告.docx": model_report(),
        "实验分析报告.docx": experiment(),
    }
    for name, doc in reports.items():
        doc.save(OUT / name)
        print(OUT / name)


if __name__ == "__main__":
    main()
